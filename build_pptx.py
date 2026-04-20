#!/usr/bin/env python3
"""
Builds GST_Presentation.pptx from the captured screenshot frames.

Structure:
  - Each HTML slide = a sequence of PPTX slides (one per frame).
  - Frames within a section auto-advance every ~FRAME_MS ms with a fast Fade.
  - The last frame of each section requires a manual click (pause point).
  - Between sections a slow Fade transition marks the boundary.
"""

import io
import json
from pathlib import Path

import numpy as np
from PIL import Image
from lxml import etree
from pptx import Presentation
from pptx.util import Emu

SCREENSHOTS_DIR = Path('/home/user/designs/screenshots')
OUTPUT_FILE     = Path('/home/user/designs/GST_Presentation.pptx')

# Slide canvas (16:9 widescreen)
SLIDE_W = Emu(12192000)   # 13.33 in × 914400
SLIDE_H = Emu(6858000)    # 7.50  in × 914400

# Frames are "similar" (skip) if fewer than this fraction of pixels differ by >PX_CHANGE_MIN
DUPE_MIN_CHANGED_FRAC = 0.0005   # 0.05% of pixels = ~1036 px on 1920×1080
DUPE_PX_CHANGE_MIN    = 12       # per-channel max change required to count a pixel as "changed"

# PPTX XML namespaces
P_NS = 'http://schemas.openxmlformats.org/presentationml/2006/main'
A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'


# ─────────────────────────────────────────────────────────────────────────────
# Low-level XML helpers
# ─────────────────────────────────────────────────────────────────────────────

def _p(tag, **attrib):
    return etree.Element(f'{{{P_NS}}}{tag}', **attrib)

def _sub(parent, ns, tag, **attrib):
    return etree.SubElement(parent, f'{{{ns}}}{tag}', **attrib)

def _psub(parent, tag, **attrib):
    return _sub(parent, P_NS, tag, **attrib)


def make_transition(fast: bool, is_last_in_section: bool, frame_ms: int) -> etree._Element:
    """
    fast=True  → quick cross-fade between animation frames.
    fast=False → slow fade at section boundaries.
    is_last_in_section=True → manual click only (no auto-advance timing).
    """
    dur = '120' if fast else '700'
    attrib = {'dur': dur, 'spd': 'fast' if fast else 'slow'}

    if not is_last_in_section:
        # auto-advance + click both work
        attrib['advClick'] = '1'
        attrib['advTm']    = str(frame_ms)
    else:
        attrib['advClick'] = '1'   # click only — presenter pauses here

    trans = _p('transition', **attrib)
    etree.SubElement(trans, f'{{{P_NS}}}fade')
    return trans


# ─────────────────────────────────────────────────────────────────────────────
# Image helpers
# ─────────────────────────────────────────────────────────────────────────────

def load_frame(path: str | Path) -> np.ndarray:
    img = Image.open(path).convert('RGB')
    if img.size != (1920, 1080):
        img = img.resize((1920, 1080), Image.LANCZOS)
    return np.array(img, dtype=np.uint8)


def frames_are_similar(a: np.ndarray, b: np.ndarray) -> bool:
    diff = np.abs(a.astype(np.int16) - b.astype(np.int16)).max(axis=2)
    changed_frac = float((diff > DUPE_PX_CHANGE_MIN).sum()) / diff.size
    return changed_frac < DUPE_MIN_CHANGED_FRAC


def to_png_bytes(arr: np.ndarray) -> bytes:
    buf = io.BytesIO()
    Image.fromarray(arr).save(buf, format='PNG', optimize=False)
    return buf.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
# PPTX slide construction
# ─────────────────────────────────────────────────────────────────────────────

def add_full_image_slide(prs: Presentation, layout, img_bytes: bytes,
                         transition_el: etree._Element) -> None:
    """Add a blank slide filled with img_bytes and attach transition_el."""
    slide = prs.slides.add_slide(layout)

    # Fill the entire slide canvas with the screenshot
    slide.shapes.add_picture(
        io.BytesIO(img_bytes),
        left=0, top=0,
        width=int(SLIDE_W), height=int(SLIDE_H),
    )

    # Attach transition to the slide XML element
    slide._element.append(transition_el)


# ─────────────────────────────────────────────────────────────────────────────
# Main build
# ─────────────────────────────────────────────────────────────────────────────

def build_pptx() -> None:
    manifest_path = SCREENSHOTS_DIR / 'manifest.json'
    if not manifest_path.exists():
        raise FileNotFoundError(
            'Run capture_slides.js first to generate screenshots/manifest.json'
        )

    with open(manifest_path) as f:
        manifest = json.load(f)

    frame_ms = manifest.get('frameMs', 300)

    prs = Presentation()
    prs.slide_width  = SLIDE_W
    prs.slide_height = SLIDE_H
    blank_layout = prs.slide_layouts[6]  # truly blank layout

    total_slides = 0

    for slide_info in manifest['slides']:
        file_name  = slide_info['file']
        frame_list = slide_info['frames']

        print(f'  {file_name}')

        # ── Deduplicate visually identical consecutive frames ─────────────────
        selected: list[dict] = []
        prev_arr: np.ndarray | None = None

        for frame in frame_list:
            fp = Path(frame['path'])
            if not fp.exists():
                continue
            arr = load_frame(fp)
            if prev_arr is None or not frames_are_similar(arr, prev_arr):
                selected.append({'arr': arr, **frame})
                prev_arr = arr

        if not selected:
            print(f'    ⚠  No usable frames found, skipping')
            continue

        print(f'    {len(selected)} unique frames (of {len(frame_list)} captured)')

        # ── Emit one PPTX slide per unique frame ──────────────────────────────
        for idx, frame in enumerate(selected):
            is_last = idx == len(selected) - 1
            is_first = idx == 0

            # First frame of a section gets slow fade (section boundary marker)
            fast = not is_first
            transition = make_transition(
                fast=fast,
                is_last_in_section=is_last,
                frame_ms=frame_ms,
            )

            add_full_image_slide(
                prs, blank_layout,
                to_png_bytes(frame['arr']),
                transition,
            )
            total_slides += 1

    prs.save(str(OUTPUT_FILE))
    print(f'\n✓  {total_slides} slides  →  {OUTPUT_FILE}')


if __name__ == '__main__':
    build_pptx()
